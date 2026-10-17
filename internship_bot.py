#!/usr/bin/env python3
"""
Production-Ready Automated Internship & Job-Hunting Bot
Single-File Architecture with Google Sheets Persistence, OpenPyXL Reporting, and SMTP Alerts.
"""

import os
import sys
import re
import json
import time
import html
import smtplib
import hashlib
import logging
import urllib.parse
from datetime import datetime, timedelta, timezone, date
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Tuple, Set, Any
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication

import requests
import urllib3
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Google Sheets & Auth
try:
    import gspread
    from google.oauth2.service_account import Credentials
    HAS_GOOGLE_SHEETS = True
except ImportError:
    HAS_GOOGLE_SHEETS = False

# Disable insecure request warnings when TimesJobs SSL verification is disabled
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ==============================================================================
# SECTION 1: CONSTANTS AND CONFIGURABLE THRESHOLDS
# ==============================================================================

IST = timezone(timedelta(hours=5, minutes=30))

# Operational Limits
EXCEL_WINDOW_DAYS: int = 50
API_QUERIES_PER_RUN: int = int(os.getenv("API_QUERIES_PER_RUN", "4"))
DETAIL_FETCH_LIMIT: int = 40
API_RECORD_CAP: int = 150
TRACKER_ADD_LIMIT: int = 10
TRACKER_SCORE_THRESHOLD: int = 60
DEADLINE_WINDOW_DAYS: int = 14
REQUEST_TIMEOUT: int = 20
APIFY_TIMEOUT: int = 150
EMAIL_MAX_PER_SECTION: int = 25
STRICT_LOCATION: bool = os.getenv("STRICT_LOCATION", "true").lower() in ("true", "1", "yes")

# Target Search Topics (15 phrases)
SEARCH_PHRASES: List[str] = [
    "finance",
    "investment banking",
    "equity research",
    "venture capital",
    "private equity",
    "strategy",
    "consulting",
    "founders office",
    "chief of staff",
    "business analyst",
    "financial analyst fresher",
    "investment analyst fresher",
    "management trainee",
    "graduate trainee",
    "graduate program"
]

# Noise Filter (Reject titles containing these words at word boundaries)
# Exception: "chief of staff" is explicitly preserved
NOISE_TERMS: List[str] = [
    "software", "developer", "devops", "data scientist", "sales", "marketing",
    "social media", "influencer", "seo", "copywriting", "brand", "recruiter",
    "hr", "telecaller", "bpo", "customer", "graphic", "content",
    "business development", "bde", "bd", "bdm", "logistics", "senior",
    "sr", "manager", "vp", "vice president", "director", "head of",
    "lead", "principal", "avp", "chief"
]

# Domain Matching (First-Match-Wins with Assigned Points)
DOMAINS_CONFIG: Dict[str, Dict[str, Any]] = {
    "VC/PE": {
        "score": 30,
        "keywords": [r"\bVC\b", r"\bPE\b", r"venture\s+capital", r"private\s+equity", r"venture\s+partner", r"early\s+stage"]
    },
    "Investment Banking": {
        "score": 30,
        "keywords": [r"investment\s+banking", r"investment\s+analyst", r"\bIB\b", r"\bM&A\b", r"mergers\s+(?:and|&)\s+acquisitions"]
    },
    "Equity Research": {
        "score": 27,
        "keywords": [r"equity\s+research", r"equity\s+analyst", r"securities\s+research"]
    },
    "Founder's Office/Chief of Staff": {
        "score": 27,
        "keywords": [
            r"founder['’]?s\s+office", r"ceo['’]?s\s+office", r"chief\s+of\s+staff",
            r"special\s+projects", r"president['’]?s\s+office"
        ]
    },
    "Consulting/Strategy": {
        "score": 25,
        "keywords": [r"consulting", r"management\s+consulting", r"\bstrategy\b", r"strategic\s+initiatives"]
    },
    "Portfolio/Asset Management": {
        "score": 25,
        "keywords": [r"portfolio\s+management", r"asset\s+management", r"\bAMC\b", r"wealth\s+management"]
    },
    "Strategy & Operations": {
        "score": 22,
        "keywords": [r"strategy\s*&\s*ops", r"strategy\s+and\s+operations", r"business\s+operations", r"program\s+management"]
    },
    "Corporate Development/IR": {
        "score": 22,
        "keywords": [r"corporate\s+development", r"corp\s*dev", r"investor\s+relations", r"\bIR\b"]
    },
    "Research & Policy": {
        "score": 18,
        "keywords": [r"research\s*&\s*policy", r"public\s+policy", r"economic\s+research"]
    },
    "General Finance": {
        "score": 15,
        "keywords": [r"finance", r"financial\s+analyst", r"corporate\s+finance", r"financial\s+modelling", r"business\s+analyst"]
    }
}

# NCR Region & Remote Locations
NCR_LOCATIONS: List[str] = [
    r"\bdelhi\b", r"\bnew\s+delhi\b", r"\bgurugram\b", r"\bgurgaon\b",
    r"\bnoida\b", r"\bgreater\s+noida\b", r"\bfaridabad\b", r"\bghaziabad\b",
    r"\bmanesar\b", r"\bsonipat\b", r"\bbahadurgarh\b", r"\bdwarka\b",
    r"\bdelhi\s*ncr\b", r"\bncr\b"
]

# Tier-1 Employer Matching (~80+ Global and Indian Institutions)
TIER_1_COMPANIES: List[str] = [
    # Global Investment Banks
    "Goldman Sachs", "Morgan Stanley", "JPMorgan", "JP Morgan", "Barclays", "Citi", "Citigroup",
    "Bank of America", "BofA", "UBS", "Deutsche Bank", "Credit Suisse", "HSBC",
    # Elite & Mid-Market Boutiques
    "Avendus", "Ambit", "Equirus", "DAM Capital", "O3 Capital", "Nuvama", "Edelweiss", "IIFL",
    "Rothschild", "Lazard", "Evercore", "Houlihan Lokey", "Jefferies", "Moelis", "Centrum", "JM Financial",
    # VC / PE Funds
    "Peak XV", "Sequoia", "Accel", "Blume", "Blume Ventures", "Elevation", "Elevation Capital",
    "3one4", "3one4 Capital", "Kedaara", "Kedaara Capital", "Lightspeed", "Matrix Partners",
    "Nexus Venture Partners", "Info Edge", "Stellaris", "Stellaris Venture", "Chiratae",
    "Antler", "Speciale Invest", "Kae Capital", "Inventus", "KKR", "Carlyle", "Bain Capital",
    "Warburg Pincus", "TPG", "General Atlantic", "Blackstone", "Tiger Global", "SoftBank",
    # Management Consulting
    "McKinsey", "BCG", "Boston Consulting Group", "Bain & Company", "Oliver Wyman",
    "Kearney", "Strategy&", "Roland Berger", "Arthur D. Little", "ADL", "L.E.K.", "LEK Consulting",
    "Alvarez & Marsal", "Deloitte", "PwC", "EY-Parthenon", "EY", "KPMG", "Redseer", "Zinnov",
    # Asset Management / Mutual Funds
    "ICICI Prudential", "SBI Funds", "Nippon India", "HDFC AMC", "Kotak AMC", "Axis MF",
    "Mirae Asset", "DSP Mutual Fund", "UTI AMC", "Aditya Birla Sun Life", "Franklin Templeton",
    "Quant Mutual Fund", "Parag Parikh", "PPFAS", "WhiteOak", "360 ONE", "Marcellus"
]

# Tracker Dropdown Allowed Values
TRACKER_STATUS_OPTIONS: List[str] = [
    "To Apply", "Applied", "Interview", "Offer", "Rejected", "Ghosted", "Withdrawn"
]

# Realistic HTTP Headers for Browser Emulation
DEFAULT_HEADERS: Dict[str, str] = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/125.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "en-IN,en-GB;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept-Encoding": "gzip, deflate",
    "Connection": "keep-alive"
}

# ==============================================================================
# SECTION 2: DATA MODELS & TELEMETRY
# ==============================================================================

@dataclass
class JobListing:
    title: str
    company: str
    location: str
    link: str
    source: str
    description: str = ""
    posted_date: str = ""
    deadline: str = ""
    score: int = 0
    role_type: str = "Internship"
    domain: str = "General Finance"
    is_tier1: bool = False
    is_ppo: bool = False
    has_entry_signal: bool = False
    score_breakdown: List[str] = field(default_factory=list)
    fingerprint: str = ""
    normalized_url: str = ""
    found_on: str = ""
    is_new: bool = True

@dataclass
class SourceDiagnostic:
    source_name: str
    source_type: str  # HTML Scraper, API, Apify
    http_status: int = 0
    host_path: str = ""
    cards_found: int = 0
    page_title: str = ""
    is_blocked: bool = False
    error_message: str = ""

# ==============================================================================
# SECTION 3: CREDENTIAL SCRUBBING & LOGGING
# ==============================================================================

class SecretScrubber:
    """Masks registered sensitive secrets from all console and log messages."""
    def __init__(self):
        self._secrets: Set[str] = set()

    def register_secret(self, secret: Optional[str]) -> None:
        if secret and len(str(secret).strip()) >= 4:
            self._secrets.add(str(secret).strip())

    def scrub(self, text: str) -> str:
        if not text:
            return ""
        scrubbed = str(text)
        for s in self._secrets:
            scrubbed = scrubbed.replace(s, "[REDACTED]")
        # Regex sanitization for standard authorization / API patterns
        scrubbed = re.sub(r'(bearer\s+)[a-zA-Z0-9_\-\.]{8,}', r'\1[REDACTED]', scrubbed, flags=re.I)
        scrubbed = re.sub(r'([?&](?:key|api_key|token|password|secret|app_key)=)[^&]+', r'\1[REDACTED]', scrubbed, flags=re.I)
        return scrubbed

class RedactingFormatter(logging.Formatter):
    def __init__(self, fmt: str, datefmt: str, scrubber: SecretScrubber):
        super().__init__(fmt, datefmt)
        self.scrubber = scrubber

    def format(self, record: logging.LogRecord) -> str:
        orig = super().format(record)
        return self.scrubber.scrub(orig)

def setup_logging(scrubber: SecretScrubber) -> logging.Logger:
    logger = logging.getLogger("InternshipBot")
    logger.setLevel(logging.INFO)
    for h in list(logger.handlers):
        logger.removeHandler(h)

    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(RedactingFormatter(
        fmt="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
        scrubber=scrubber
    ))
    logger.addHandler(handler)
    return logger

# Global security instance
SCRUBBER = SecretScrubber()
logger = setup_logging(SCRUBBER)

# ==============================================================================
# SECTION 4: NORMALIZATION, PARSING & SECURITY HELPERS
# ==============================================================================

def normalize_url(url: str) -> str:
    """Removes tracking query parameters, hashes, and trailing slashes."""
    if not url:
        return ""
    try:
        parsed = urllib.parse.urlparse(url.strip())
        query_dict = urllib.parse.parse_qs(parsed.query)
        # Drop marketing / tracking queries
        filtered_query = {
            k: v for k, v in query_dict.items()
            if not k.lower().startswith("utm_") and k.lower() not in (
                "ref", "src", "trk", "fbclid", "gclid", "campaign", "source", "medium"
            )
        }
        clean_path = parsed.path.rstrip("/")
        normalized = urllib.parse.urlunparse((
            parsed.scheme.lower(),
            parsed.netloc.lower(),
            clean_path,
            "",
            urllib.parse.urlencode(filtered_query, doseq=True),
            ""
        ))
        return normalized
    except Exception:
        return url.strip()

def compute_fingerprint(company: str, title: str) -> str:
    """Generates a deterministic 16-character SHA-256 fingerprint for deduplication."""
    clean_company = re.sub(r'[^a-z0-9]', '', (company or '').lower())
    clean_title = re.sub(r'[^a-z0-9]', '', (title or '').lower())
    combined = f"{clean_company}|{clean_title}".encode('utf-8')
    return hashlib.sha256(combined).hexdigest()[:16]

def parse_date_robust(date_input: Any) -> Optional[str]:
    """
    Parses dates from string formats, relative expressions, or Google Sheets serial numbers.
    Returns ISO date string (YYYY-MM-DD) or None.
    """
    if date_input is None:
        return None

    # Handle float / int Google Sheets or Excel serial numbers (epoch 1899-12-30)
    if isinstance(date_input, (int, float)):
        try:
            parsed_dt = datetime(1899, 12, 30) + timedelta(days=float(date_input))
            return parsed_dt.strftime("%Y-%m-%d")
        except Exception:
            return None

    date_str = str(date_input).strip()
    if not date_str:
        return None

    # Check numeric 5-digit Excel serial passed as string (e.g. "45601")
    if date_str.isdigit() and len(date_str) == 5:
        try:
            parsed_dt = datetime(1899, 12, 30) + timedelta(days=float(date_str))
            return parsed_dt.strftime("%Y-%m-%d")
        except Exception:
            pass

    date_lower = date_str.lower()
    now_ist = datetime.now(IST)

    # Relative expressions
    if "today" in date_lower or "just now" in date_lower:
        return now_ist.strftime("%Y-%m-%d")
    if "yesterday" in date_lower:
        return (now_ist - timedelta(days=1)).strftime("%Y-%m-%d")

    days_ago_match = re.search(r'(\d+)\s+(?:day|days)\s+ago', date_lower)
    if days_ago_match:
        try:
            return (now_ist - timedelta(days=int(days_ago_match.group(1)))).strftime("%Y-%m-%d")
        except Exception:
            pass

    hours_ago_match = re.search(r'(\d+)\s+(?:hour|hours|hr|hrs)\s+ago', date_lower)
    if hours_ago_match:
        return now_ist.strftime("%Y-%m-%d")

    # Common standard textual date patterns (e.g., 2026-10-17, 17/10/2026, 17 Oct 2026)
    formats = [
        "%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y",
        "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
        "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%SZ"
    ]
    # Clean ordinal suffixes (1st, 2nd, 3rd, 4th)
    clean_text = re.sub(r'(\d+)(st|nd|rd|th)', r'\1', date_str)
    for fmt in formats:
        try:
            return datetime.strptime(clean_text, fmt).strftime("%Y-%m-%d")
        except Exception:
            continue

    return None

def extract_deadline_from_text(text: str) -> Optional[str]:
    """Extracts application deadlines matching 'apply by', 'deadline', or 'last date'."""
    if not text:
        return None
    now_date = datetime.now(IST).date()
    patterns = [
        r'(?:last\s+date|apply\s+by|deadline|closes\s+on)[\s:]*([0-9]{1,2}(?:st|nd|rd|th)?\s+[a-zA-Z]{3,9}\s*[0-9]{2,4})',
        r'(?:last\s+date|apply\s+by|deadline|closes\s+on)[\s:]*([0-9]{4}[-/][0-9]{1,2}[-/][0-9]{1,2})',
        r'(?:last\s+date|apply\s+by|deadline|closes\s+on)[\s:]*([0-9]{1,2}[-/][0-9]{1,2}[-/][0-9]{2,4})'
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            candidate = parse_date_robust(match.group(1))
            if candidate:
                try:
                    candidate_date = datetime.strptime(candidate, "%Y-%m-%d").date()
                    if candidate_date >= now_date:
                        return candidate
                except Exception:
                    continue
    return None

def sanitize_for_spreadsheet(value: Any) -> str:
    """Defends against formula injection in Google Sheets and Excel."""
    s = str(value or "")
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{s}"
    return s

def validate_url(url: str) -> bool:
    """Ensures URLs use safe HTTP or HTTPS protocols and avoid javascript: attacks."""
    if not url:
        return False
    parsed = urllib.parse.urlparse(url)
    return parsed.scheme.lower() in ("http", "https") and bool(parsed.netloc)

def get_host_path_only(url: str) -> str:
    """Extracts host and path only, completely stripping sensitive query parameters."""
    if not url:
        return ""
    try:
        parsed = urllib.parse.urlparse(url)
        return f"{parsed.scheme}://{parsed.netloc}{parsed.path}"
    except Exception:
        return "invalid-url"

# ==============================================================================
# SECTION 5: FILTERING & SCORING RULES
# ==============================================================================

def passes_title_gate(title: str) -> bool:
    """
    Rejects irrelevant or senior role titles.
    Strictly preserves 'Chief of Staff' while rejecting other 'Chief' roles.
    """
    if not title:
        return False
    title_clean = title.strip().lower()

    # Exception rule: Allow "Chief of Staff"
    if re.search(r'\bchief\s+of\s+staff\b', title_clean):
        return True

    for noise in NOISE_TERMS:
        if noise == "chief":
            if re.search(r'\bchief\b', title_clean) and not re.search(r'\bchief\s+of\s+staff\b', title_clean):
                return False
        else:
            if re.search(rf'\b{re.escape(noise)}\b', title_clean):
                return False

    return True

def check_credential_rule(text: str) -> bool:
    """
    Rejects listings explicitly demanding MBA/PGDM only, CA Inter/Final, or CFA Level requirements,
    UNLESS undergraduate degrees (BBA, B.Com, Bachelor, Undergraduate, Graduate) are also accepted.
    Note: 'Postgraduate' is not treated as 'graduate'.
    """
    if not text:
        return True
    text_lower = text.lower()

    has_prohibitive = bool(re.search(
        r'\b(?:mba/pgdm\s+only|pursuing\s+mba|pursuing\s+pgdm|pursuing\s+cfa|ca\s+inter|ca\s+final|cfa\s+level)\b',
        text_lower
    ))
    if not has_prohibitive:
        return True

    # Check for undergraduate exceptions
    text_clean = text_lower.replace("postgraduate", "").replace("post-graduate", "")
    has_undergrad_exception = bool(re.search(
        r'\b(?:bba|b\.com|bcom|bachelor|undergraduate|graduate)\b',
        text_clean
    ))
    return has_undergrad_exception

def check_experience_rule(text: str, role_type: str) -> bool:
    """
    Experience Rules:
    - Internships requiring >= 3 years are rejected.
    - Full-time roles requiring > 2 years are rejected.
    - Ambiguous full-time roles without explicit fresher / entry / 0-2 yrs indication are rejected.
    """
    text_lower = (text or "").lower()

    # Extract years requirement (e.g. '2-5 yrs', 'minimum 3 years', '3+ years of experience')
    exp_matches = re.findall(r'(\d+)\s*(?:-|to)\s*(\d+)\s*(?:yrs|years)', text_lower)
    min_exp = None
    max_exp = None

    if exp_matches:
        try:
            min_exp = int(exp_matches[0][0])
            max_exp = int(exp_matches[0][1])
        except Exception:
            pass
    else:
        single_match = re.search(r'(?:min|minimum|at\s+least)?\s*(\d+)\+?\s*(?:yrs|years)(?:\s+of\s+experience|\s+exp)?', text_lower)
        if single_match:
            try:
                min_exp = int(single_match.group(1))
            except Exception:
                pass

    if "internship" in role_type.lower():
        if (min_exp is not None and min_exp >= 3) or (max_exp is not None and max_exp >= 3):
            return False
        return True

    # Full-time roles
    if min_exp is not None and min_exp > 2:
        return False
    if max_exp is not None and max_exp > 2:
        return False

    # For full-time, accept if explicit entry-level signal is present
    has_entry_signal = bool(re.search(
        r'\b(?:fresher|freshers|entry[- ]level|0[- ]1\s*(?:yrs|years)|0[- ]2\s*(?:yrs|years)|campus|2027\s*batch|recent\s*graduate)\b',
        text_lower
    ))
    return has_entry_signal

def classify_role_type(title: str, text: str) -> Tuple[str, int]:
    """
    Classifies listing into:
    - PPO Internship (20 pts)
    - Graduate Program (18 pts)
    - Full-time (Fresher) (14 pts)
    - Internship (12 pts)
    """
    combined = f"{title} {text}".lower()

    if re.search(r'\b(?:ppo|pre[- ]placement\s+offer|conversion\s+opportunity)\b', combined):
        return "PPO Internship", 20
    if re.search(r'\b(?:graduate\s+program|graduate\s+trainee|management\s+trainee|campus\s+recruitment|2027\s+batch)\b', combined):
        return "Graduate Program", 18
    if re.search(r'\b(?:full[- ]time|permanent)\b', combined):
        return "Full-time (Fresher)", 14

    return "Internship", 12

def match_domain(text: str) -> Optional[Tuple[str, int]]:
    """Evaluates 10 domain categories using First-Match-Wins with assigned points."""
    text_clean = text.lower()
    for domain_name, config in DOMAINS_CONFIG.items():
        for kw in config["keywords"]:
            if re.search(kw, text_clean, re.I):
                return domain_name, config["score"]
    return None

def is_ncr_or_remote(location_str: str) -> Tuple[bool, bool, int]:
    """
    Checks if location matches Delhi NCR (+10 pts) or Remote (+8 pts).
    Returns (is_ncr, is_remote, location_points).
    """
    if not location_str:
        return False, False, 0
    loc_clean = location_str.lower()

    for ncr_pattern in NCR_LOCATIONS:
        if re.search(ncr_pattern, loc_clean, re.I):
            return True, False, 10

    if re.search(r'\b(?:remote|work\s+from\s+home|wfh|anywhere)\b', loc_clean, re.I):
        return False, True, 8

    return False, False, 0

def check_tier1_company(company_str: str) -> Tuple[bool, int]:
    """Matches company name against Tier-1 roster using boundary-safe lookarounds."""
    if not company_str:
        return False, 0
    comp = company_str.strip()
    for t1 in TIER_1_COMPANIES:
        # Safe boundary match supporting symbols like Strategy&
        pattern = rf"(?<![a-zA-Z0-9]){re.escape(t1)}(?![a-zA-Z0-9])"
        if re.search(pattern, comp, re.I):
            return True, 30
    return False, 0

def score_job_listing(listing: JobListing) -> None:
    """
    Calculates final score (max 100) using multi-factor evaluation:
    - Tier-1 Company: 30 pts
    - Domain Fit: up to 30 pts
    - Role Type: 12-20 pts
    - Location: NCR (10) or Remote (8)
    - Entry-Level Signal: 6 pts
    - Deadline <= 14 days: 4 pts
    """
    total = 0
    breakdown = []
    combined_text = f"{listing.title} {listing.company} {listing.location} {listing.description}"

    # 1. Tier-1 Match (30 pts)
    is_t1, t1_score = check_tier1_company(listing.company)
    if is_t1:
        total += t1_score
        listing.is_tier1 = True
        breakdown.append(f"Tier-1 Employer (+{t1_score})")

    # 2. Domain Match (up to 30 pts)
    domain_match = match_domain(combined_text)
    if domain_match:
        dom_name, dom_score = domain_match
        listing.domain = dom_name
        total += dom_score
        breakdown.append(f"{dom_name} (+{dom_score})")
    else:
        listing.domain = "General Finance"
        total += 15
        breakdown.append("General Finance (+15)")

    # 3. Role Type (12 - 20 pts)
    role_type, role_score = classify_role_type(listing.title, combined_text)
    listing.role_type = role_type
    listing.is_ppo = (role_type == "PPO Internship")
    total += role_score
    breakdown.append(f"{role_type} (+{role_score})")

    # 4. Location Bonus (NCR 10 pts, Remote 8 pts)
    is_ncr, is_remote, loc_score = is_ncr_or_remote(listing.location)
    if is_ncr:
        total += loc_score
        breakdown.append("Delhi-NCR (+10)")
    elif is_remote:
        total += loc_score
        breakdown.append("Remote (+8)")

    # 5. Entry-Level Signal (6 pts)
    if re.search(r'\b(?:fresher|entry[- ]level|0[- ]1\s*yrs?|campus|2027\s*batch|recent\s*graduate)\b', combined_text, re.I):
        total += 6
        listing.has_entry_signal = True
        breakdown.append("Entry-Level Signal (+6)")

    # 6. Deadline Urgency (4 pts if deadline is within 14 days)
    if listing.deadline:
        try:
            today_date = datetime.now(IST).date()
            dl_date = datetime.strptime(listing.deadline, "%Y-%m-%d").date()
            days_left = (dl_date - today_date).days
            if 0 <= days_left <= DEADLINE_WINDOW_DAYS:
                total += 4
                breakdown.append(f"Urgent Deadline <= 14d (+4, {days_left}d left)")
        except Exception:
            pass

    listing.score = min(total, 100)
    listing.score_breakdown = breakdown

# ==============================================================================
# SECTION 6: WEB SCRAPING & API ADAPTERS
# ==============================================================================

class ScraperSession:
    """Manages HTTP session, retries, block detection, and diagnostics."""
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update(DEFAULT_HEADERS)
        self.diagnostics: List[SourceDiagnostic] = []

    def request(self, method: str, url: str, verify: bool = True, **kwargs) -> Tuple[Optional[requests.Response], SourceDiagnostic]:
        diag = SourceDiagnostic(
            source_name=kwargs.pop("source_name", "HTTP"),
            source_type=kwargs.pop("source_type", "HTML Scraper"),
            host_path=get_host_path_only(url)
        )
        kwargs.setdefault("timeout", REQUEST_TIMEOUT)

        for attempt in range(2):
            try:
                resp = self.session.request(method, url, verify=verify, **kwargs)
                diag.http_status = resp.status_code

                # Detect Cloudflare or block pages (only on HTTP 403/429 or explicit challenge signatures)
                if resp.status_code in (403, 429) or (
                    resp.status_code != 200 and any(
                        sig in resp.text.lower()
                        for sig in ("cloudflare", "access denied", "security challenge", "just a moment...")
                    )
                ) or (
                    "<title>just a moment...</title>" in resp.text.lower() or
                    "<title>access denied</title>" in resp.text.lower()
                ):
                    diag.is_blocked = True
                    diag.error_message = f"Block page detected (HTTP {resp.status_code})"

                # Capture page title if zero cards
                try:
                    soup = BeautifulSoup(resp.text[:5000], "html.parser")
                    if soup.title and soup.title.string:
                        diag.page_title = soup.title.string.strip()[:100]
                except Exception:
                    pass

                return resp, diag

            except Exception as e:
                diag.error_message = SCRUBBER.scrub(str(e))
                if attempt == 0:
                    time.sleep(3)
                else:
                    return None, diag

        return None, diag

def scrape_shine(session: ScraperSession, query: str) -> List[JobListing]:
    """Scrapes Shine.com using verified article and result-card CSS selectors."""
    listings: List[JobListing] = []
    slug = urllib.parse.quote_plus(query.replace(" ", "-"))
    url = f"https://www.shine.com/job-search/{slug}-jobs"

    resp, diag = session.request("GET", url, source_name="Shine", source_type="HTML Scraper")
    if not resp or resp.status_code != 200:
        session.diagnostics.append(diag)
        return listings

    soup = BeautifulSoup(resp.text, "html.parser")
    # Verified modern Shine card selectors
    cards = (
        soup.select("article[class*='result-card']") or
        soup.find_all("article") or
        soup.select("div.jobCard") or
        soup.select("div[class*='job_card']") or
        soup.select("div[class*='srp-job']")
    )
    diag.cards_found = len(cards)
    session.diagnostics.append(diag)

    for card in cards:
        try:
            title_el = (
                card.select_one("h3[class*='role'] a") or
                card.select_one("a[class*='hit']") or
                card.select_one("a.job_title_anchor") or
                card.select_one("h2 > a") or
                card.select_one("h3 > a") or
                card.find("a")
            )
            comp_el = (
                card.select_one("span[class*='company']") or
                card.select_one("div[class*='co'] span") or
                card.select_one("span.company_name") or
                card.select_one("div[class*='company']")
            )
            loc_el = (
                card.select_one("span[class*='meta-text']") or
                card.select_one("span.loc") or
                card.select_one("span[class*='location']")
            )

            if not title_el or not comp_el:
                continue

            title = title_el.get_text(strip=True)
            company = comp_el.get_text(strip=True)
            location = loc_el.get_text(strip=True) if loc_el else "Delhi NCR"
            href = title_el.get("href", "")
            full_url = urllib.parse.urljoin("https://www.shine.com", href)

            listings.append(JobListing(
                title=title,
                company=company,
                location=location,
                link=full_url,
                source="Shine"
            ))
        except Exception:
            continue

    return listings

def scrape_timesjobs(session: ScraperSession, query: str) -> List[JobListing]:
    """
    Scrapes TimesJobs.
    Note: SSL verification is disabled (verify=False) due to TimesJobs' known broken certificate chain.
    """
    listings: List[JobListing] = []
    url = (
        f"https://www.timesjobs.com/candidate/job-search.html"
        f"?searchType=personal498&from=submit&txtKeywords={urllib.parse.quote_plus(query)}&txtLocation="
    )

    resp, diag = session.request("GET", url, verify=False, source_name="TimesJobs", source_type="HTML Scraper")
    if not resp or resp.status_code != 200:
        session.diagnostics.append(diag)
        return listings

    soup = BeautifulSoup(resp.text, "html.parser")
    cards = soup.select("li.clearfix.job-bx") or soup.select("div.job-container") or soup.select("div[class*='srp-listing']")
    diag.cards_found = len(cards)
    session.diagnostics.append(diag)

    for card in cards:
        try:
            title_el = card.select_one("h2 > a") or card.select_one("h3 > a")
            comp_el = card.select_one("h3.joblist-comp-name") or card.select_one("span.joblist-comp-name")
            loc_el = card.select_one("ul.top-jd-dtl li span") or card.select_one("span[title]")

            if not title_el or not comp_el:
                continue

            title = title_el.get_text(strip=True)
            company = re.sub(r'\s+', ' ', comp_el.get_text(strip=True))
            location = loc_el.get_text(strip=True) if loc_el else "Delhi NCR"
            full_url = title_el.get("href", "")

            listings.append(JobListing(
                title=title,
                company=company,
                location=location,
                link=full_url,
                source="TimesJobs"
            ))
        except Exception:
            continue

    return listings

def scrape_internshala(session: ScraperSession, query: str) -> List[JobListing]:
    """Scrapes Internshala (best-effort, graceful degradation on anti-bot challenge)."""
    listings: List[JobListing] = []
    slug = urllib.parse.quote_plus(query.replace(" ", "-"))
    url = f"https://internshala.com/internships/keywords-{slug}"

    resp, diag = session.request("GET", url, source_name="Internshala", source_type="HTML Scraper")
    if not resp or resp.status_code != 200:
        session.diagnostics.append(diag)
        return listings

    soup = BeautifulSoup(resp.text, "html.parser")
    cards = soup.select("div.individual_internship") or soup.select("div[class*='internship_meta']")
    diag.cards_found = len(cards)
    session.diagnostics.append(diag)

    for card in cards:
        try:
            title_el = (
                card.select_one("a.job-title-href") or
                card.select_one("h2.job-internship-name a") or
                card.select_one("a[id='job_title']") or
                card.select_one("a.view_detail_button") or
                card.select_one("h3 > a")
            )
            comp_el = (
                card.select_one("p.company-name") or
                card.select_one("p.company_name") or
                card.select_one("div.company_name") or
                card.select_one("a.link_display_like_text")
            )
            loc_el = (
                card.select_one("div.row-1-item.locations span a") or
                card.select_one("div.row-1-item.locations span") or
                card.select_one("div.locations") or
                card.select_one("a#location_names") or
                card.select_one("span.location_link")
            )
            desc_el = card.select_one("div.about_job div.text") or card.select_one("div.text")

            if not title_el or not comp_el:
                continue

            title = title_el.get_text(strip=True)
            company = comp_el.get_text(strip=True)
            location = loc_el.get_text(strip=True) if loc_el else "Delhi NCR"
            href = title_el.get("href") or card.get("data-href") or ""
            full_url = urllib.parse.urljoin("https://internshala.com", href)
            description = desc_el.get_text(strip=True) if desc_el else ""

            listings.append(JobListing(
                title=title,
                company=company,
                location=location,
                link=full_url,
                source="Internshala",
                description=description
            ))
        except Exception:
            continue

    return listings

# ----------------- Optional APIs (Enabled Only When Secret Exists) -----------------

def fetch_adzuna(session: ScraperSession, query: str) -> List[JobListing]:
    """Queries Adzuna API (restricted to India) if secrets exist."""
    app_id = os.getenv("ADZUNA_APP_ID")
    app_key = os.getenv("ADZUNA_APP_KEY")
    if not app_id or not app_key:
        return []

    SCRUBBER.register_secret(app_id)
    SCRUBBER.register_secret(app_key)

    url = (
        f"https://api.adzuna.com/v1/api/jobs/in/search/1"
        f"?app_id={app_id}&app_key={app_key}&results_per_page=50&what={urllib.parse.quote_plus(query)}"
        f"&where=India&content-type=application/json"
    )

    resp, diag = session.request("GET", url, source_name="Adzuna", source_type="API")
    session.diagnostics.append(diag)
    listings: List[JobListing] = []

    if resp and resp.status_code == 200:
        data = resp.json().get("results", [])
        diag.cards_found = len(data)
        for item in data:
            title = item.get("title", "")
            company = item.get("company", {}).get("display_name", "")
            location = item.get("location", {}).get("display_name", "India")
            link = item.get("redirect_url", "")
            desc = item.get("description", "")
            if title and company and link:
                listings.append(JobListing(
                    title=title, company=company, location=location,
                    link=link, source="Adzuna", description=desc
                ))
    return listings

def fetch_jsearch(session: ScraperSession, query: str) -> List[JobListing]:
    """Queries JSearch API via RapidAPI if RAPIDAPI_KEY exists."""
    rapidapi_key = os.getenv("RAPIDAPI_KEY")
    if not rapidapi_key:
        return []
    SCRUBBER.register_secret(rapidapi_key)

    url = "https://jsearch.p.rapidapi.com/search"
    headers = {
        **DEFAULT_HEADERS,
        "X-RapidAPI-Key": rapidapi_key,
        "X-RapidAPI-Host": "jsearch.p.rapidapi.com"
    }
    params = {"query": f"{query} in India", "page": "1", "num_pages": "1", "date_posted": "week"}

    resp, diag = session.request("GET", url, headers=headers, params=params, source_name="JSearch", source_type="API")
    session.diagnostics.append(diag)
    listings: List[JobListing] = []

    if resp and resp.status_code == 200:
        data = resp.json().get("data", [])
        diag.cards_found = len(data)
        for item in data:
            title = item.get("job_title", "")
            company = item.get("employer_name", "")
            location = item.get("job_city") or item.get("job_country") or "India"
            link = item.get("job_apply_link") or item.get("job_google_link", "")
            desc = item.get("job_description", "")
            if title and company and link:
                listings.append(JobListing(
                    title=title, company=company, location=location,
                    link=link, source="JSearch", description=desc
                ))
    return listings

def fetch_jooble(session: ScraperSession, query: str) -> List[JobListing]:
    """Queries Jooble API if JOOBLE_KEY exists."""
    jooble_key = os.getenv("JOOBLE_KEY")
    if not jooble_key:
        return []
    SCRUBBER.register_secret(jooble_key)

    url = f"https://jooble.org/api/{jooble_key}"
    payload = {"keywords": query, "location": "India", "page": 1}

    resp, diag = session.request("POST", url, json=payload, source_name="Jooble", source_type="API")
    session.diagnostics.append(diag)
    listings: List[JobListing] = []

    if resp and resp.status_code == 200:
        data = resp.json().get("jobs", [])
        diag.cards_found = len(data)
        for item in data:
            title = item.get("title", "")
            company = item.get("company", "")
            location = item.get("location", "India")
            link = item.get("link", "")
            snippet = item.get("snippet", "")
            if title and company and link:
                listings.append(JobListing(
                    title=title, company=company, location=location,
                    link=link, source="Jooble", description=snippet
                ))
    return listings

def fetch_serpapi(session: ScraperSession, query: str) -> List[JobListing]:
    """Queries SerpAPI Google Jobs if SERPAPI_KEY exists."""
    serpapi_key = os.getenv("SERPAPI_KEY")
    if not serpapi_key:
        return []
    SCRUBBER.register_secret(serpapi_key)

    url = "https://serpapi.com/search.json"
    params = {"engine": "google_jobs", "q": query, "location": "India", "api_key": serpapi_key}

    resp, diag = session.request("GET", url, params=params, source_name="SerpAPI", source_type="API")
    session.diagnostics.append(diag)
    listings: List[JobListing] = []

    if resp and resp.status_code == 200:
        data = resp.json().get("jobs_results", [])
        diag.cards_found = len(data)
        for item in data:
            title = item.get("title", "")
            company = item.get("company_name", "")
            location = item.get("location", "India")
            link = ""
            related = item.get("related_links", [])
            if related and isinstance(related, list) and "link" in related[0]:
                link = related[0]["link"]
            elif item.get("share_link"):
                link = item.get("share_link")

            desc = item.get("description", "")
            if title and company and link:
                listings.append(JobListing(
                    title=title, company=company, location=location,
                    link=link, source="SerpAPI", description=desc
                ))
    return listings

def fetch_apify(session: ScraperSession, query: str) -> List[JobListing]:
    """Queries configured Apify actors via REST API with 150s timeout."""
    token = os.getenv("APIFY_TOKEN")
    actors_str = os.getenv("APIFY_ACTORS", "")
    if not token or not actors_str.strip():
        return []
    SCRUBBER.register_secret(token)

    actors = [a.strip() for a in actors_str.split(",") if a.strip()]
    raw_template = os.getenv("APIFY_INPUT_JSON", '{"queries": ["{query}"], "maxItems": 50}')
    max_items = int(os.getenv("APIFY_MAX_ITEMS", "100"))

    listings: List[JobListing] = []
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}

    for actor in actors:
        # Safe JSON parameter replacement
        try:
            formatted_json_str = raw_template.replace("{query}", query).replace("{location}", "India")
            payload = json.loads(formatted_json_str)
        except Exception as e:
            logger.warning(f"Apify template JSON formatting failed: {e}")
            continue

        url = f"https://api.apify.com/v2/acts/{actor}/run-sync-get-dataset-items?timeout={APIFY_TIMEOUT}&memory=1024"
        diag = SourceDiagnostic(source_name=f"Apify ({actor})", source_type="Apify", host_path=get_host_path_only(url))

        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=APIFY_TIMEOUT)
            diag.http_status = resp.status_code
            if resp.status_code == 201 or resp.status_code == 200:
                items = resp.json()
                if isinstance(items, list):
                    diag.cards_found = len(items)
                    for item in items[:max_items]:
                        title = item.get("title") or item.get("jobTitle") or item.get("name")
                        company = item.get("company") or item.get("companyName") or item.get("employer_name")
                        link = item.get("url") or item.get("link") or item.get("jobUrl") or item.get("applyUrl")
                        loc = item.get("location") or item.get("city") or "India"
                        desc = item.get("description") or item.get("snippet") or ""

                        if title and company and link:
                            listings.append(JobListing(
                                title=str(title), company=str(company), location=str(loc),
                                link=str(link), source=f"Apify-{actor}", description=str(desc)
                            ))
            else:
                diag.error_message = f"Apify returned HTTP {resp.status_code}"
        except Exception as e:
            diag.error_message = SCRUBBER.scrub(str(e))

        session.diagnostics.append(diag)

    return listings

# ==============================================================================
# SECTION 7: GOOGLE SHEETS PERSISTENCE (LIVE SPREADSHEET BACKEND)
# ==============================================================================

class GoogleSheetsManager:
    """Manages connection, tab initialization, cross-run deduplication, and tracker sync."""
    def __init__(self, sheet_id: Optional[str], service_account_json: Optional[str]):
        self.sheet_id = sheet_id
        self.service_account_json = service_account_json
        self.client = None
        self.spreadsheet = None
        self.is_connected = False

    def connect(self) -> bool:
        if not HAS_GOOGLE_SHEETS or not self.sheet_id or not self.service_account_json:
            logger.info("Google Sheets credentials not provided; proceeding in local memory mode.")
            return False

        try:
            SCRUBBER.register_secret(self.sheet_id)
            info = json.loads(self.service_account_json.strip())
            scopes = [
                "https://www.googleapis.com/auth/spreadsheets",
                "https://www.googleapis.com/auth/drive"
            ]
            creds = Credentials.from_service_account_info(info, scopes=scopes)
            self.client = gspread.authorize(creds)
            self.spreadsheet = self.client.open_by_key(self.sheet_id)
            self.is_connected = True
            logger.info("Successfully authenticated with Google Sheets.")
            self._ensure_tabs()
            return True
        except Exception as e:
            logger.warning(f"Failed to connect to Google Sheets: {SCRUBBER.scrub(str(e))}")
            self.is_connected = False
            return False

    def _ensure_tabs(self) -> None:
        """Ensures seen_listings, all_listings, and tracker tabs exist with correct headers."""
        if not self.is_connected or not self.spreadsheet:
            return

        expected_tabs = {
            "seen_listings": ["url", "fingerprint", "status", "seen_on"],
            "all_listings": [
                "title", "company", "location", "link", "source", "score",
                "type", "domain", "tier1", "ppo", "deadline", "found_on"
            ],
            "tracker": [
                "Company", "Role", "Type", "Link", "Score", "Status",
                "Applied Date", "Deadline", "Follow-up Date", "Contact",
                "PPO", "Notes", "Added On"
            ]
        }

        existing_sheets = {ws.title: ws for ws in self.spreadsheet.worksheets()}
        for tab_name, headers in expected_tabs.items():
            if tab_name not in existing_sheets:
                try:
                    ws = self.spreadsheet.add_worksheet(title=tab_name, rows=1000, cols=len(headers) + 2)
                    ws.append_row(headers)
                    logger.info(f"Created Google Sheets tab: {tab_name}")
                except Exception as e:
                    logger.warning(f"Could not create tab {tab_name}: {e}")

    def load_seen_registry(self) -> Set[str]:
        """Loads seen URLs and fingerprints from seen_listings tab for cross-run deduplication."""
        seen = set()
        if not self.is_connected or not self.spreadsheet:
            return seen

        try:
            ws = self.spreadsheet.worksheet("seen_listings")
            records = ws.get_all_values()
            if len(records) > 1:
                for row in records[1:]:
                    if len(row) >= 2:
                        url_val = row[0].strip()
                        fp_val = row[1].strip()
                        if url_val:
                            seen.add(url_val)
                        if fp_val:
                            seen.add(fp_val)
            logger.info(f"Loaded {len(seen)} existing keys from Google Sheets seen_listings.")
        except Exception as e:
            logger.warning(f"Error loading seen_listings: {e}")

        return seen

    def append_seen_listings(self, entries: List[Tuple[str, str, str, str]]) -> None:
        """Batch appends seen entries (url, fingerprint, status, seen_on)."""
        if not self.is_connected or not self.spreadsheet or not entries:
            return
        try:
            ws = self.spreadsheet.worksheet("seen_listings")
            ws.append_rows(entries)
            logger.info(f"Appended {len(entries)} records to seen_listings.")
        except Exception as e:
            logger.warning(f"Error updating seen_listings: {e}")

    def append_all_listings(self, listings: List[JobListing], today_str: str) -> None:
        """Batch appends qualified listings to all_listings master archive (no description column)."""
        if not self.is_connected or not self.spreadsheet or not listings:
            return
        try:
            ws = self.spreadsheet.worksheet("all_listings")
            rows = []
            for item in listings:
                rows.append([
                    sanitize_for_spreadsheet(item.title),
                    sanitize_for_spreadsheet(item.company),
                    sanitize_for_spreadsheet(item.location),
                    item.link,
                    item.source,
                    item.score,
                    item.role_type,
                    item.domain,
                    "Yes" if item.is_tier1 else "No",
                    "Yes" if item.is_ppo else "No",
                    item.deadline or "",
                    today_str
                ])
            ws.append_rows(rows)
            logger.info(f"Appended {len(rows)} listings to all_listings master archive.")
        except Exception as e:
            logger.warning(f"Error updating all_listings: {e}")

    def append_tracker_listings(self, listings: List[JobListing], today_str: str) -> None:
        """Appends top listings (score >= 60, max 10) to tracker tab."""
        if not self.is_connected or not self.spreadsheet:
            return

        qualifying = [l for l in listings if l.score >= TRACKER_SCORE_THRESHOLD]
        qualifying.sort(key=lambda x: x.score, reverse=True)
        to_add = qualifying[:TRACKER_ADD_LIMIT]

        if not to_add:
            return

        try:
            ws = self.spreadsheet.worksheet("tracker")
            rows = []
            for item in to_add:
                rows.append([
                    sanitize_for_spreadsheet(item.company),
                    sanitize_for_spreadsheet(item.title),
                    item.role_type,
                    item.link,
                    item.score,
                    "To Apply",  # Default status
                    "",          # Applied Date
                    item.deadline or "",
                    "",          # Follow-up Date
                    "",          # Contact
                    "Yes" if item.is_ppo else "No",
                    "",          # Notes
                    today_str
                ])
            ws.append_rows(rows)
            logger.info(f"Added {len(rows)} top opportunities to application tracker tab.")
        except Exception as e:
            logger.warning(f"Error updating tracker tab: {e}")

    def fetch_50day_history(self, cutoff_date_str: str) -> List[Dict[str, Any]]:
        """Retrieves listings from the last 50 days from all_listings tab."""
        if not self.is_connected or not self.spreadsheet:
            return []
        try:
            ws = self.spreadsheet.worksheet("all_listings")
            records = ws.get_all_records()
            filtered = []
            for r in records:
                found_on = str(r.get("found_on", ""))
                if found_on and found_on >= cutoff_date_str:
                    filtered.append(r)
            return filtered
        except Exception as e:
            logger.warning(f"Error retrieving 50-day history from Google Sheets: {e}")
            return []

# ==============================================================================
# SECTION 8: OPENPYXL EXCEL DIGEST GENERATION
# ==============================================================================

def generate_excel_report(
    output_path: str,
    today_listings: List[JobListing],
    historical_listings: List[Dict[str, Any]],
    diagnostics: List[SourceDiagnostic],
    today_str: str
) -> bool:
    """
    Generates internships.xlsx with two sheets:
    1. 'All Listings' (50-day window, freeze header, autofilter, navy style)
    2. 'Source Stats' (Telemetry table)
    """
    try:
        wb = openpyxl.Workbook()

        # Sheet 1: All Listings
        ws1 = wb.active
        ws1.title = "All Listings"

        headers1 = [
            "New Today", "Score", "Company", "Role", "Type", "Domain",
            "Location", "Tier 1", "PPO", "Deadline", "Apply Link", "Source", "Found On"
        ]
        ws1.append(headers1)

        # Build list of rows
        rows_data = []

        # Add today's newly scraped listings
        for l in today_listings:
            rows_data.append({
                "new": "NEW",
                "score": l.score,
                "company": l.company,
                "role": l.title,
                "type": l.role_type,
                "domain": l.domain,
                "location": l.location,
                "tier1": "Yes" if l.is_tier1 else "No",
                "ppo": "Yes" if l.is_ppo else "No",
                "deadline": l.deadline or "N/A",
                "link": l.link,
                "source": l.source,
                "found_on": today_str
            })

        # Add historical items from Google Sheets (excluding today to avoid duplicates)
        for h in historical_listings:
            h_found = str(h.get("found_on", ""))
            if h_found != today_str:
                rows_data.append({
                    "new": "",
                    "score": int(h.get("score", 0)),
                    "company": str(h.get("company", "")),
                    "role": str(h.get("title", "")),
                    "type": str(h.get("type", "")),
                    "domain": str(h.get("domain", "")),
                    "location": str(h.get("location", "")),
                    "tier1": str(h.get("tier1", "No")),
                    "ppo": str(h.get("ppo", "No")),
                    "deadline": str(h.get("deadline", "N/A")),
                    "link": str(h.get("link", "")),
                    "source": str(h.get("source", "")),
                    "found_on": h_found
                })

        # Sort: found_on descending, then score descending
        rows_data.sort(key=lambda x: (x["found_on"], x["score"]), reverse=True)

        for row in rows_data:
            ws1.append([
                row["new"], row["score"], row["company"], row["role"],
                row["type"], row["domain"], row["location"], row["tier1"],
                row["ppo"], row["deadline"], row["link"], row["source"], row["found_on"]
            ])

        # Style Sheet 1
        navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        new_badge_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        new_badge_font = Font(name="Calibri", size=10, bold=True, color="166534")

        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=1, column=col_idx)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx in range(2, ws1.max_row + 1):
            for col_idx in range(1, len(headers1) + 1):
                cell = ws1.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

            # Style "NEW" badge column
            new_cell = ws1.cell(row=row_idx, column=1)
            if new_cell.value == "NEW":
                new_cell.fill = new_badge_fill
                new_cell.font = new_badge_font
                new_cell.alignment = Alignment(horizontal="center", vertical="center")

            # Make link clickable
            link_cell = ws1.cell(row=row_idx, column=11)
            if link_cell.value and validate_url(str(link_cell.value)):
                link_cell.hyperlink = str(link_cell.value)
                link_cell.font = Font(color="2563EB", underline="single")

        ws1.freeze_panes = "A2"
        ws1.auto_filter.ref = ws1.dimensions

        # Auto-fit column widths
        for col in ws1.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Sheet 2: Source Stats
        ws2 = wb.create_sheet(title="Source Stats")
        headers2 = ["Source Name", "Source Type", "HTTP Status", "Cards Found", "Blocked?", "Diagnostics / Errors"]
        ws2.append(headers2)

        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for diag in diagnostics:
            ws2.append([
                diag.source_name,
                diag.source_type,
                diag.http_status if diag.http_status else "N/A",
                diag.cards_found,
                "YES" if diag.is_blocked else "No",
                diag.error_message or "OK"
            ])

        for row_idx in range(2, ws2.max_row + 1):
            for col_idx in range(1, len(headers2) + 1):
                cell = ws2.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        ws2.freeze_panes = "A2"
        for col in ws2.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

        wb.save(output_path)
        logger.info(f"Generated Excel workbook at: {output_path}")
        return True
    except Exception as e:
        logger.error(f"Failed to generate Excel report: {e}", exc_info=True)
        return False

# ==============================================================================
# SECTION 9: EMAIL DISPATCH & HEALTH ALERTS (GMAIL SMTP SSL / TLS)
# ==============================================================================

def build_card_html(listing: JobListing) -> str:
    """Builds a responsive HTML card for a single job listing."""
    badges_html = ""
    if listing.is_tier1:
        badges_html += '<span style="background-color: #FEF3C7; color: #92400E; font-size: 11px; font-weight: bold; padding: 2px 8px; border-radius: 4px; margin-right: 6px;">Tier-1</span>'
    if listing.is_ppo:
        badges_html += '<span style="background-color: #E0E7FF; color: #3730A3; font-size: 11px; font-weight: bold; padding: 2px 8px; border-radius: 4px; margin-right: 6px;">PPO</span>'

    safe_title = html.escape(listing.title)
    safe_company = html.escape(listing.company)
    safe_location = html.escape(listing.location)
    safe_domain = html.escape(listing.domain)
    safe_type = html.escape(listing.role_type)
    safe_source = html.escape(listing.source)
    deadline_text = f"• Deadline: <strong>{html.escape(listing.deadline)}</strong>" if listing.deadline else ""

    apply_btn = (
        f'<a href="{html.escape(listing.link)}" style="display: inline-block; background-color: #1E3A8A; color: #FFFFFF; text-decoration: none; font-size: 13px; font-weight: bold; padding: 6px 14px; border-radius: 6px; margin-top: 8px;">Apply Now →</a>'
        if validate_url(listing.link) else '<span style="color: #94A3B8; font-size: 12px;">Link unavailable</span>'
    )

    card = f"""
    <div style="background-color: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 8px; padding: 14px; margin-bottom: 12px; box-shadow: 0 1px 2px rgba(0,0,0,0.05);">
        <div style="display: flex; justify-content: space-between; align-items: flex-start;">
            <div>
                <h3 style="margin: 0 0 4px 0; font-size: 16px; color: #0F172A;">{safe_title}</h3>
                <div style="font-size: 13px; color: #475569; margin-bottom: 6px;">
                    <strong>{safe_company}</strong> • {safe_location} • Source: {safe_source} {deadline_text}
                </div>
            </div>
            <div style="text-align: right;">
                <span style="display: inline-block; background-color: #DCFCE7; color: #166534; font-size: 14px; font-weight: bold; padding: 4px 10px; border-radius: 12px;">Score: {listing.score}</span>
            </div>
        </div>
        <div style="margin-top: 4px; font-size: 12px; color: #64748B;">
            {badges_html}
            <span style="background-color: #F1F5F9; color: #334155; padding: 2px 8px; border-radius: 4px; margin-right: 6px;">{safe_type}</span>
            <span style="background-color: #F8FAFC; color: #475569; border: 1px solid #CBD5E1; padding: 2px 8px; border-radius: 4px;">{safe_domain}</span>
        </div>
        <div>
            {apply_btn}
        </div>
    </div>
    """
    return card

def send_daily_email(
    today_listings: List[JobListing],
    excel_path: str,
    sheet_id: Optional[str] = None
) -> bool:
    """Dispatches the daily 8:00 AM IST HTML summary with internships.xlsx attached."""
    sender = os.getenv("EMAIL")
    password = os.getenv("PASSWORD", "").replace(" ", "")
    recipient = os.getenv("EMAIL_TO") or sender

    if not sender or not password:
        logger.warning("SMTP credentials (EMAIL / PASSWORD) not set. Skipping email dispatch.")
        return False

    SCRUBBER.register_secret(password)
    SCRUBBER.register_secret(sender)

    now_ist_str = datetime.now(IST).strftime("%d %b %Y")
    subject = f"🚀 Daily Job & Internship Digest — {now_ist_str} ({len(today_listings)} New)"

    # Group listings into 3 sections
    internships = [l for l in today_listings if "internship" in l.role_type.lower()]
    grad_programs = [l for l in today_listings if "graduate" in l.role_type.lower() or "trainee" in l.role_type.lower()]
    full_time = [l for l in today_listings if "full-time" in l.role_type.lower()]

    internships.sort(key=lambda x: x.score, reverse=True)
    grad_programs.sort(key=lambda x: x.score, reverse=True)
    full_time.sort(key=lambda x: x.score, reverse=True)

    def render_section(title: str, items: List[JobListing]) -> str:
        if not items:
            return f"<h2 style='color: #1E3A8A; border-bottom: 2px solid #E2E8F0; padding-bottom: 6px;'>{title} (0)</h2><p style='color: #64748B;'>No new openings in this category today.</p>"

        cards_html = "".join([build_card_html(item) for item in items[:EMAIL_MAX_PER_SECTION]])
        overflow_note = ""
        if len(items) > EMAIL_MAX_PER_SECTION:
            overflow_note = f"<p style='color: #2563EB; font-weight: bold;'>+{len(items) - EMAIL_MAX_PER_SECTION} more in attached internships.xlsx</p>"

        return f"""
        <h2 style='color: #1E3A8A; border-bottom: 2px solid #E2E8F0; padding-bottom: 6px;'>{title} ({len(items)})</h2>
        {cards_html}
        {overflow_note}
        """

    sheet_link_banner = ""
    if sheet_id:
        sheet_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit"
        sheet_link_banner = f"""
        <div style="background-color: #EFF6FF; border: 1px solid #BFDBFE; border-radius: 8px; padding: 12px; margin-bottom: 20px; text-align: center;">
            <p style="margin: 0; font-size: 14px; color: #1E40AF;">
                📊 <strong>Live Google Sheet:</strong> All past 50-day records and your personal Application Tracker are updated live.<br>
                <a href="{sheet_url}" style="color: #1D4ED8; font-weight: bold; text-decoration: underline;">Open Live Google Spreadsheet →</a>
            </p>
        </div>
        """

    body_html = f"""
    <!DOCTYPE html>
    <html>
    <head><meta charset="utf-8"></head>
    <body style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; background-color: #F8FAFC; padding: 20px; margin: 0;">
        <div style="max-width: 720px; margin: 0 auto; background-color: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 12px; overflow: hidden; padding: 24px;">
            <div style="background: linear-gradient(135deg, #1E3A8A, #3B82F6); padding: 20px; border-radius: 8px; color: #FFFFFF; margin-bottom: 20px;">
                <h1 style="margin: 0; font-size: 22px;">Daily Job & Internship Digest</h1>
                <p style="margin: 6px 0 0 0; font-size: 14px; opacity: 0.9;">Curated Finance, Strategy, IB & VC Opportunities — {now_ist_str}</p>
            </div>

            {sheet_link_banner}

            {render_section("1. Internships & PPO Roles", internships)}
            {render_section("2. Graduate Programs & Trainee Roles", grad_programs)}
            {render_section("3. Full-Time Roles (Freshers)", full_time)}

            <div style="margin-top: 30px; padding-top: 15px; border-top: 1px solid #E2E8F0; font-size: 12px; color: #94A3B8; text-align: center;">
                Automated Job-Hunting Bot • Preserved in attached internships.xlsx (50-Day Archive)
            </div>
        </div>
    </body>
    </html>
    """

    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = recipient

    msg.attach(MIMEText(body_html, "html"))

    # Attach Excel file
    if os.path.exists(excel_path):
        try:
            with open(excel_path, "rb") as f:
                part = MIMEApplication(f.read(), _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                part.add_header("Content-Disposition", "attachment", filename="internships.xlsx")
                msg.attach(part)
        except Exception as e:
            logger.error(f"Failed to attach Excel file to email: {e}")

    # Send via SMTP SSL (port 465) or STARTTLS (port 587)
    try:
        smtp_port = int(os.getenv("SMTP_PORT", "465"))
        if smtp_port == 465:
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=30) as server:
                server.login(sender, password)
                server.send_message(msg)
        else:
            with smtplib.SMTP("smtp.gmail.com", smtp_port, timeout=30) as server:
                server.starttls()
                server.login(sender, password)
                server.send_message(msg)

        logger.info(f"Daily digest email successfully sent to: {recipient}")
        return True
    except Exception as e:
        logger.error(f"Failed to send daily digest email: {SCRUBBER.scrub(str(e))}")
        return False

def send_alert_email(diagnostics: List[SourceDiagnostic], excel_path: str) -> bool:
    """Dispatches a diagnostic ALERT email when all enabled sources return zero listings."""
    sender = os.getenv("EMAIL")
    password = os.getenv("PASSWORD", "").replace(" ", "")
    recipient = os.getenv("EMAIL_TO") or sender

    if not sender or not password:
        logger.warning("SMTP credentials missing; cannot send alert email.")
        return False

    now_ist_str = datetime.now(IST).strftime("%d %b %Y %H:%M:%S")
    subject = f"⚠️ [ALERT] Job Bot Zero-Results Diagnostic — {now_ist_str} IST"

    rows_html = ""
    for d in diagnostics:
        status_badge = f'<span style="color: {"red" if d.is_blocked else "green"}; font-weight: bold;">{"BLOCKED" if d.is_blocked else "OK"}</span>'
        rows_html += f"""
        <tr>
            <td style="padding: 8px; border: 1px solid #CBD5E1;">{html.escape(d.source_name)}</td>
            <td style="padding: 8px; border: 1px solid #CBD5E1;">{html.escape(d.source_type)}</td>
            <td style="padding: 8px; border: 1px solid #CBD5E1;">{d.http_status or 'N/A'}</td>
            <td style="padding: 8px; border: 1px solid #CBD5E1;">{d.cards_found}</td>
            <td style="padding: 8px; border: 1px solid #CBD5E1;">{status_badge}</td>
            <td style="padding: 8px; border: 1px solid #CBD5E1; font-size: 11px;">{html.escape(d.page_title or 'N/A')}</td>
            <td style="padding: 8px; border: 1px solid #CBD5E1; font-size: 11px; color: #DC2626;">{html.escape(d.error_message or 'None')}</td>
        </tr>
        """

    body_html = f"""
    <html>
    <body style="font-family: sans-serif; padding: 20px;">
        <div style="max-width: 800px; margin: 0 auto; border: 1px solid #EF4444; border-radius: 8px; padding: 20px;">
            <h2 style="color: #DC2626; margin-top: 0;">⚠️ Job Scraper Alert: 0 Listings Found Across All Sources</h2>
            <p>Every scraper and API adapter returned 0 valid job cards during today's run at <strong>{now_ist_str} IST</strong>.</p>

            <h3>Per-Source Diagnostic Telemetry</h3>
            <table style="width: 100%; border-collapse: collapse; font-size: 13px;">
                <thead>
                    <tr style="background-color: #F1F5F9;">
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Source</th>
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Type</th>
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Status</th>
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Cards</th>
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Blocked?</th>
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Page Title</th>
                        <th style="padding: 8px; border: 1px solid #CBD5E1;">Errors</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                </tbody>
            </table>

            <p style="margin-top: 20px; font-size: 13px; color: #475569;">
                The 50-day historical archive from your Google Sheet has still been compiled and attached as <code>internships.xlsx</code>.
            </p>
        </div>
    </body>
    </html>
    """

    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = recipient
    msg.attach(MIMEText(body_html, "html"))

    if os.path.exists(excel_path):
        try:
            with open(excel_path, "rb") as f:
                part = MIMEApplication(f.read(), _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                part.add_header("Content-Disposition", "attachment", filename="internships.xlsx")
                msg.attach(part)
        except Exception:
            pass

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=30) as server:
            server.login(sender, password)
            server.send_message(msg)
        logger.info("Diagnostic health alert email dispatched.")
        return True
    except Exception as e:
        logger.error(f"Failed to dispatch alert email: {SCRUBBER.scrub(str(e))}")
        return False

# ==============================================================================
# SECTION 10: MAIN ORCHESTRATION PIPELINE
# ==============================================================================

def run_pipeline() -> int:
    """Executes the complete production pipeline."""
    start_time = time.time()
    today_dt = datetime.now(IST)
    today_str = today_dt.strftime("%Y-%m-%d")
    logger.info(f"=== Starting Production Job-Hunting Run: {today_dt.strftime('%Y-%m-%d %H:%M:%S')} IST ===")

    # Register known environment variables in Scrubber
    for env_key in [
        "EMAIL", "PASSWORD", "GSHEET_ID", "GOOGLE_SERVICE_ACCOUNT_JSON",
        "ADZUNA_APP_ID", "ADZUNA_APP_KEY", "RAPIDAPI_KEY", "JOOBLE_KEY",
        "SERPAPI_KEY", "APIFY_TOKEN"
    ]:
        SCRUBBER.register_secret(os.getenv(env_key))

    # 1. Initialize Google Sheets Persistence (Option A)
    gsheet_mgr = GoogleSheetsManager(
        sheet_id=os.getenv("GSHEET_ID"),
        service_account_json=os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON")
    )
    gsheet_mgr.connect()

    # Load seen registry for cross-run deduplication
    seen_registry: Set[str] = gsheet_mgr.load_seen_registry()
    in_run_seen: Set[str] = set()

    # 2. Select Search Phrases (15 for HTML scrapers, 4 rotated for APIs)
    day_of_year = today_dt.timetuple().tm_yday
    start_idx = (day_of_year * API_QUERIES_PER_RUN) % len(SEARCH_PHRASES)
    rotated_api_queries = [
        SEARCH_PHRASES[(start_idx + i) % len(SEARCH_PHRASES)]
        for i in range(API_QUERIES_PER_RUN)
    ]
    logger.info(f"Executing with {len(SEARCH_PHRASES)} HTML search phrases and {len(rotated_api_queries)} rotated API phrases: {rotated_api_queries}")

    session = ScraperSession()
    raw_listings: List[JobListing] = []

    # 3. Scrape HTML Sources (All 15 Phrases)
    for phrase in SEARCH_PHRASES:
        try:
            raw_listings.extend(scrape_shine(session, phrase))
        except Exception as e:
            logger.warning(f"Shine error on '{phrase}': {SCRUBBER.scrub(str(e))}")

        try:
            raw_listings.extend(scrape_timesjobs(session, phrase))
        except Exception as e:
            logger.warning(f"TimesJobs error on '{phrase}': {SCRUBBER.scrub(str(e))}")

        try:
            raw_listings.extend(scrape_internshala(session, phrase))
        except Exception as e:
            logger.warning(f"Internshala error on '{phrase}': {SCRUBBER.scrub(str(e))}")

    # 4. Scrape APIs & Apify (4 Rotated Phrases to protect quotas)
    for phrase in rotated_api_queries:
        try:
            raw_listings.extend(fetch_adzuna(session, phrase))
            raw_listings.extend(fetch_jsearch(session, phrase))
            raw_listings.extend(fetch_jooble(session, phrase))
            raw_listings.extend(fetch_serpapi(session, phrase))
            raw_listings.extend(fetch_apify(session, phrase))
        except Exception as e:
            logger.warning(f"API scraping error on '{phrase}': {SCRUBBER.scrub(str(e))}")

    logger.info(f"Total raw listings scraped from all sources: {len(raw_listings)}")

    # 5. Pipeline Deduplication, Filtering & Scoring
    new_seen_entries: List[Tuple[str, str, str, str]] = []
    candidates_to_score: List[JobListing] = []

    for item in raw_listings:
        norm_url = normalize_url(item.link)
        fp = compute_fingerprint(item.company, item.title)
        item.normalized_url = norm_url
        item.fingerprint = fp

        # In-run and Cross-run Deduplication
        if norm_url in seen_registry or fp in seen_registry or norm_url in in_run_seen or fp in in_run_seen:
            continue

        in_run_seen.add(norm_url)
        in_run_seen.add(fp)

        # Title Gate (Reject noise terms early)
        if not passes_title_gate(item.title):
            new_seen_entries.append((norm_url, fp, "rejected_title", today_str))
            continue

        candidates_to_score.append(item)

    logger.info(f"Passed initial title gate & dedup: {len(candidates_to_score)} listings.")

    # 6. Pre-scoring & Detail-Page Fetching
    # Sort candidates by pre-score so the 40 detail-fetch slots go to top prospects
    for item in candidates_to_score:
        is_t1, t1_pts = check_tier1_company(item.company)
        is_ncr, is_rem, loc_pts = is_ncr_or_remote(item.location)
        dom_match = match_domain(item.title)
        dom_pts = dom_match[1] if dom_match else 15
        item.score = t1_pts + loc_pts + dom_pts

    candidates_to_score.sort(key=lambda x: x.score, reverse=True)

    # Fetch details for top 40 listings needing description
    detail_fetches_done = 0
    api_records_used = 0

    qualified_today: List[JobListing] = []

    for item in candidates_to_score:
        # If record already has rich description from API/Apify, skip HTTP detail fetch
        if item.description and len(item.description) > 50:
            if api_records_used < API_RECORD_CAP:
                api_records_used += 1
            else:
                continue
        elif detail_fetches_done < DETAIL_FETCH_LIMIT and validate_url(item.link):
            try:
                resp, _ = session.request("GET", item.link, source_name="DetailFetch")
                if resp and resp.status_code == 200:
                    soup = BeautifulSoup(resp.text[:30000], "html.parser")
                    # Extract body text
                    for s in soup(["script", "style", "nav", "footer"]):
                        s.decompose()
                    desc_text = " ".join(soup.stripped_strings)
                    item.description = desc_text[:2000]
                    # Attempt deadline extraction
                    extracted_dl = extract_deadline_from_text(desc_text)
                    if extracted_dl:
                        item.deadline = extracted_dl
                detail_fetches_done += 1
            except Exception:
                pass

        # Quality Checks (Credential & Experience Rules)
        full_text = f"{item.title} {item.company} {item.location} {item.description}"

        if not check_credential_rule(full_text):
            new_seen_entries.append((item.normalized_url, item.fingerprint, "rejected_credentials", today_str))
            continue

        if not check_experience_rule(full_text, item.role_type):
            new_seen_entries.append((item.normalized_url, item.fingerprint, "rejected_experience", today_str))
            continue

        # Strict Location Enforcement
        if STRICT_LOCATION:
            is_ncr, is_rem, _ = is_ncr_or_remote(item.location)
            if not is_ncr and not is_rem:
                new_seen_entries.append((item.normalized_url, item.fingerprint, "rejected_location", today_str))
                continue

        # Final Scoring & Classification
        score_job_listing(item)
        item.found_on = today_str
        qualified_today.append(item)
        new_seen_entries.append((item.normalized_url, item.fingerprint, "kept", today_str))

    logger.info(f"Final scored & qualified listings for today: {len(qualified_today)}")

    # 7. Update Live Google Sheet
    gsheet_mgr.append_seen_listings(new_seen_entries)
    gsheet_mgr.append_all_listings(qualified_today, today_str)
    gsheet_mgr.append_tracker_listings(qualified_today, today_str)

    # 8. Generate Excel Report (internships.xlsx) with 50-day window
    cutoff_50d = (today_dt - timedelta(days=EXCEL_WINDOW_DAYS)).strftime("%Y-%m-%d")
    historical_50d = gsheet_mgr.fetch_50day_history(cutoff_50d)

    excel_output_path = "internships.xlsx"
    generate_excel_report(
        output_path=excel_output_path,
        today_listings=qualified_today,
        historical_listings=historical_50d,
        diagnostics=session.diagnostics,
        today_str=today_str
    )

    # 9. Email Notifications
    all_zero = (len(raw_listings) == 0)
    if all_zero:
        logger.warning("All sources returned zero listings. Triggering diagnostic alert email...")
        send_alert_email(session.diagnostics, excel_output_path)
    else:
        send_daily_email(qualified_today, excel_output_path, sheet_id=os.getenv("GSHEET_ID"))

    elapsed = round(time.time() - start_time, 2)
    logger.info(f"=== Execution Complete in {elapsed}s. {len(qualified_today)} opportunities recorded. ===")
    return 0

if __name__ == "__main__":
    sys.exit(run_pipeline())
